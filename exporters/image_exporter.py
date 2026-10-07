"""Image-description exports (spec §7.11): CSV, JSON, and a DOCX list.

Decorative status appears exactly as decided: a suggestion is only a
suggestion until a person confirmed it, and the exports say which.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path


def _decorative_text(item: dict) -> str:
    decorative = item.get("decorative") or {}
    if decorative.get("confirmed") is True:
        return "decorative (confirmed by a person)"
    if decorative.get("confirmed") is False:
        return "not decorative (confirmed by a person)"
    if decorative.get("suggested"):
        return f"suggested decorative — NOT confirmed ({decorative.get('reason', '')})"
    return "informative"


def image_provenance_lines(project: dict) -> list[str]:
    items = project.get("images", [])
    approved = sum(1 for i in items if i.get("status") == "approved")
    reviewed = items and approved == len(items)
    return [
        "Image descriptions: drafted by a vision model on this computer; "
        + ("reviewed by a person." if reviewed else "NOT yet reviewed by a person."),
        "Cloud services used: none.",
        f"Status: {'reviewed' if reviewed else 'DRAFT. Not yet reviewed by a person.'}",
    ]


def export_images_csv(project: dict, output_path: str | Path) -> Path:
    path = Path(output_path)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["# " + line for line in image_provenance_lines(project)])
        writer.writerow(
            ["name", "kind", "alt_text", "long_description", "decorative", "status", "approved_by"]
        )
        for item in project.get("images", []):
            writer.writerow(
                [
                    item.get("name", ""),
                    item.get("kind", ""),
                    item.get("alt", ""),
                    item.get("long_description", ""),
                    _decorative_text(item),
                    item.get("status", "draft"),
                    item.get("approved_by") or "",
                ]
            )
    return path


def export_images_json(project: dict, output_path: str | Path) -> Path:
    path = Path(output_path)
    path.write_text(
        json.dumps(
            {
                "provenance": image_provenance_lines(project),
                "images": project.get("images", []),
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def export_images_docx(project: dict, output_path: str | Path) -> Path:
    import docx
    from docx.shared import Pt

    path = Path(output_path)
    document = docx.Document()
    document.add_heading(
        f"Image descriptions — {project.get('title') or 'Untitled'}", level=1
    )
    for line in image_provenance_lines(project):
        p = document.add_paragraph(line)
        p.runs[0].font.size = Pt(9)

    for item in project.get("images", []):
        document.add_heading(item.get("name", "Image"), level=2)
        p = document.add_paragraph()
        p.add_run("Alt text: ").bold = True
        p.add_run(item.get("alt") or "(none)")
        if item.get("long_description"):
            p = document.add_paragraph()
            p.add_run("Long description: ").bold = True
            p.add_run(item["long_description"])
        p = document.add_paragraph()
        p.add_run("Decorative: ").bold = True
        p.add_run(_decorative_text(item))
        status = item.get("status", "draft")
        p = document.add_paragraph()
        p.add_run("Status: ").bold = True
        p.add_run(
            f"approved by {item.get('approved_by')}" if status == "approved" else "draft"
        )
    document.save(str(path))
    return path
